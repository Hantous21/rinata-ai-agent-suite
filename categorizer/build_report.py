"""
Stage 6 — Word Document Management Report
Generates Rinata_Management_Report.docx using Claude to write the narrative.
"""

import pandas as pd
import anthropic
import json
from pathlib import Path
from dotenv import dotenv_values
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from collections import Counter

# ── Paths ─────────────────────────────────────────────────────────────────────
ROOT   = Path(__file__).parent.parent
CAT    = ROOT / "output" / "categorized_reviews.csv"
RAW    = ROOT / "output" / "raw_reviews.csv"
OUTPUT = ROOT / "output" / "Rinata_Management_Report.docx"

# ── Load data ─────────────────────────────────────────────────────────────────
print("=== Rinata Word Report Builder ===\n", flush=True)
cat = pd.read_csv(CAT, encoding="utf-8")
raw = pd.read_csv(RAW, encoding="utf-8")
cat["rating"] = pd.to_numeric(cat["rating"], errors="coerce")
raw["rating"] = pd.to_numeric(raw["rating"], errors="coerce")

# ── Key stats ─────────────────────────────────────────────────────────────────
total          = len(raw)
avg_rating     = round(raw["rating"].mean(), 2)
pct_positive   = round((cat["sentiment"] == "positive").sum() / len(cat) * 100, 1)
category_counts = cat["primary_category"].value_counts().to_dict()
high_urg       = cat[cat["urgency"] == "high"]
med_urg        = cat[cat["urgency"] == "medium"]
neg_reviews    = cat[cat["sentiment"] == "negative"].sort_values("rating")

DAY_ORDER = ["Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday"]
day_stats = []
for day in DAY_ORDER:
    day_raw = raw[raw["day_of_week"] == day]
    day_cat = cat[cat["day_of_week"] == day]
    day_stats.append({
        "day":      day,
        "count":    len(day_raw),
        "avg":      round(float(day_raw["rating"].mean()), 2) if len(day_raw) > 0 else 0,
        "neg":      int((day_cat["sentiment"] == "negative").sum()),
        "top_cat":  day_cat["primary_category"].value_counts().index[0] if len(day_cat) > 0 else "N/A",
    })

# Sample quotes for top categories
def get_quotes(category, sentiment="positive", n=2):
    subset = cat[
        (cat["primary_category"] == category) &
        (cat["sentiment"] == sentiment) &
        (cat["review_text"].str.len() > 40)
    ].head(n)
    return [str(r["review_text"])[:200] for _, r in subset.iterrows()]

# ── Build data summary for Claude ────────────────────────────────────────────
data_summary = f"""
RINATA RESTAURANT — GOOGLE REVIEW ANALYSIS SUMMARY
Total reviews: {total}
Average star rating: {avg_rating} / 5.0
Reviews with written text analyzed: {len(cat)}
Percentage positive sentiment: {pct_positive}%

CATEGORY BREAKDOWN (primary category, reviews with text only):
{json.dumps(category_counts, indent=2)}

DAY OF WEEK STATS:
{json.dumps(day_stats, indent=2)}

HIGH URGENCY REVIEWS ({len(high_urg)} total):
{chr(10).join([f'- {r["rating"]}* | {r["primary_category"]} | {r["one_line_summary"]}' for _, r in high_urg.iterrows()])}

MEDIUM URGENCY REVIEWS ({len(med_urg)} total — top 8):
{chr(10).join([f'- {r["rating"]}* | {r["primary_category"]} | {r["one_line_summary"]}' for _, r in med_urg.head(8).iterrows()])}

NEGATIVE REVIEWS SAMPLE (top 6 by lowest rating):
{chr(10).join([f'- {r["rating"]}* | {r["primary_category"]} | {r["review_text"][:150]}' for _, r in neg_reviews.head(6).iterrows()])}
"""

# ── Ask Claude to write the narrative ────────────────────────────────────────
print("Asking Claude to write the narrative analysis...", flush=True)
env    = dotenv_values(ROOT / ".env")
client = anthropic.Anthropic(api_key=env["ANTHROPIC_API_KEY"])

prompt = f"""
You are an internal business analyst writing a concise management memo for the owners of Rinata,
an Italian restaurant in Minneapolis. You have analyzed their Google Reviews using AI.

Write the following sections based on the data below. Use a direct, internal memo style.
No fluff. Be specific with numbers. Use bullet points where appropriate.

Write exactly these sections with these headers (use ### for each header):

### What the Reviews Tell Us
3-4 bullet points summarizing the overall picture — what customers experience, what they say,
how Rinata compares to a typical restaurant. Use specific numbers.

### What Is Working Well
3-4 bullet points on clear strengths. Reference specific categories and percentages.
Pick 1-2 direct customer quotes that illustrate these strengths.

### What Needs Attention
Bullet points covering all high-urgency and medium-urgency themes.
Group related issues. Be direct about what the problem is and how many reviews mention it.
Do not soften the language — management needs to know exactly what is happening.

### Day-of-Week Patterns
What the day-of-week data reveals. Call out Wednesday specifically.
Note which days have the most negative reviews and what categories those tend to be.
Suggest what this might mean operationally.

### Recommended Actions
A numbered action list — specific, concrete, prioritized.
Start with the highest-urgency items. Each action should say WHO should do WHAT and WHY.
Aim for 6-8 actions.

---

DATA:
{data_summary}
"""

response = client.messages.create(
    model      = "claude-haiku-4-5-20251001",
    max_tokens = 1800,
    messages = [{"role": "user", "content": prompt}]
)
narrative = response.content[0].text.strip()
print("Narrative complete.\n", flush=True)

# ── Helper: set paragraph style ───────────────────────────────────────────────
def set_color(run, hex_color):
    r, g, b = int(hex_color[0:2],16), int(hex_color[2:4],16), int(hex_color[4:6],16)
    run.font.color.rgb = RGBColor(r, g, b)

def add_heading(doc, text, level=1, color="0F3460"):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_color(run, color)
        run.font.name = "Calibri"
    return p

def add_body(doc, text, bold=False, color=None, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.bold = bold
    if color:
        set_color(run, color)
    return p

def add_kpi_row(doc, items):
    """items = list of (label, value) tuples"""
    table = doc.add_table(rows=2, cols=len(items))
    table.style = "Table Grid"
    for col, (label, value) in enumerate(items):
        label_cell = table.cell(0, col)
        value_cell = table.cell(1, col)

        lp = label_cell.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lr = lp.add_run(label.upper())
        lr.font.name = "Calibri"
        lr.font.size = Pt(8)
        lr.bold = True
        set_color(lr, "888888")

        vp = value_cell.paragraphs[0]
        vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        vr = vp.add_run(str(value))
        vr.font.name = "Calibri"
        vr.font.size = Pt(18)
        vr.bold = True
        set_color(vr, "0F3460")

        # shade header row
        tc_pr = label_cell._tc.get_or_add_tcPr()
        shd   = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F5F5F5")
        shd.set(qn("w:val"),  "clear")
        tc_pr.append(shd)
    return table

def parse_and_write_narrative(doc, text):
    """Parse Claude's ### sections and write them into the doc."""
    lines = text.split("\n")
    for line in lines:
        line = line.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:], level=2, color="0F3460")
        elif line.startswith("**") and line.endswith("**"):
            add_body(doc, line.strip("*"), bold=True)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(line[2:])
            run.font.name = "Calibri"
            run.font.size = Pt(11)
        elif line.startswith(tuple("123456789")) and ". " in line[:4]:
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(line[line.index(". ")+2:])
            run.font.name = "Calibri"
            run.font.size = Pt(11)
        elif line.startswith("> ") or (line.startswith('"') and line.endswith('"')):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent  = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            run = p.add_run(line.lstrip("> "))
            run.font.name    = "Calibri"
            run.font.size    = Pt(10)
            run.font.italic  = True
            set_color(run, "555555")
        else:
            add_body(doc, line)

# ── Build the document ────────────────────────────────────────────────────────
doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1.2)

# Cover / Title
doc.add_paragraph()
title = doc.add_heading("Rinata Restaurant", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.name = "Calibri"
    set_color(run, "0F3460")

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Google Review Analysis — Internal Management Report")
sr.font.name = "Calibri"
sr.font.size = Pt(14)
set_color(sr, "555555")

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
dr = date_p.add_run("May 2026  |  Confidential")
dr.font.name = "Calibri"
dr.font.size = Pt(10)
set_color(dr, "AAAAAA")

doc.add_paragraph()
doc.add_paragraph()

# KPI bar
add_heading(doc, "At a Glance", level=1, color="0F3460")
add_kpi_row(doc, [
    ("Total Reviews",   total),
    ("Avg Rating",      f"{avg_rating} / 5.0"),
    ("% Positive",      f"{pct_positive}%"),
    ("High Urgency",    len(high_urg)),
    ("Top Category",    "Food Quality"),
])
doc.add_paragraph()

# How to read this report
add_heading(doc, "How to Use This Report", level=1, color="0F3460")
p = doc.add_paragraph()
run = p.add_run(
    "This report was generated by analyzing all Google Reviews for Rinata using AI. "
    "Each of the 388 written reviews was read and classified by category, sentiment, and urgency. "
    "The findings below reflect patterns across the full dataset — not individual opinions. "
    "The accompanying Excel file (Rinata_Review_Analysis.xlsx) contains the complete data "
    "and can be filtered and explored in detail."
)
run.font.name = "Calibri"
run.font.size = Pt(11)
doc.add_paragraph()

# Main narrative (Claude-written sections)
parse_and_write_narrative(doc, narrative)

doc.add_paragraph()

# Sample positive quotes
add_heading(doc, "Sample Customer Voices", level=1, color="0F3460")
add_body(doc, "Food Quality", bold=True)
for q in get_quotes("Food Quality", "positive", 2):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    run = p.add_run(f'"{q}"')
    run.font.name   = "Calibri"
    run.font.size   = Pt(10)
    run.font.italic = True
    set_color(run, "444444")

add_body(doc, "Atmosphere", bold=True)
for q in get_quotes("Atmosphere", "positive", 2):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.4)
    run = p.add_run(f'"{q}"')
    run.font.name   = "Calibri"
    run.font.size   = Pt(10)
    run.font.italic = True
    set_color(run, "444444")

doc.add_paragraph()

# Footer note
add_heading(doc, "Methodology Note", level=1, color="888888")
p = doc.add_paragraph()
run = p.add_run(
    f"Data source: Google Maps public reviews for Rinata Restaurant, Minneapolis. "
    f"Scraped May 2026. {total} total reviews collected; {len(cat)} contained written text "
    f"and were classified using Claude AI (Anthropic). "
    f"Date approximations are derived from relative timestamps (e.g. '3 weeks ago'). "
    f"Category assignments reflect AI interpretation and should be validated against the "
    f"raw data in the accompanying Excel file when making significant operational decisions."
)
run.font.name  = "Calibri"
run.font.size  = Pt(9)
run.font.italic = True
set_color(run, "888888")

# ── Save ─────────────────────────────────────────────────────────────────────
doc.save(OUTPUT)
print(f"Saved: {OUTPUT}", flush=True)
print("Done!", flush=True)
